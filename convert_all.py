import os
import re
import sys
import json
import logging
import pdfplumber
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ========================
#  Portable Base Directory
# ========================
if getattr(sys, 'frozen', False):
    BASE_DIR = os.path.dirname(sys.executable)  # for .exe
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ========================
#  Logging Setup
# ========================
logging.basicConfig(
    level=logging.INFO,
    format="%(levelname)s: %(message)s"
)

# ========================
#  Load Config
# ========================
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")
if not os.path.exists(CONFIG_PATH):
    logging.error(f"Config file not found: {CONFIG_PATH}")
    sys.exit(1)

with open(CONFIG_PATH, "r", encoding="utf-8") as f:
    config = json.load(f)

PDF_FOLDER = os.path.join(BASE_DIR, config.get("PDF_FOLDER", "invoices"))
OUTPUT_FILE = os.path.join(BASE_DIR, config.get("OUTPUT_FILE", "combined_invoices.xlsx"))
ORDER_DATE_DEFAULT = config.get("ORDER_DATE_DEFAULT", "")
SKIP_FREE_GIFTS = config.get("SKIP_FREE_GIFTS", True)

# ========================
#  Data Storage
# ========================
all_rows = []

# ========================
#  Process PDFs
# ========================
for filename in os.listdir(PDF_FOLDER):
    if filename.lower().endswith(".pdf"):
        pdf_path = os.path.join(PDF_FOLDER, filename)
        logging.info(f"🔍 Processing: {filename}")

        try:
            with pdfplumber.open(pdf_path) as pdf:
                page = pdf.pages[0]
                text = page.extract_text() or ""
                tables = page.extract_tables()

            # ==== Invoice Metadata ====
            series_match = re.search(r'Ký hiệu\s*\(Series\):\s*(\S+)', text)
            series = series_match.group(1) if series_match else ""

            number_match = re.search(r'(\d{6,8})\s*Số\s*\(No\.\)', text)
            number = number_match.group(1) if number_match else ""

            invoice_number = f"{number}.{series}" if number and series else ""

            date_match = re.search(r'Ngày\s*\(Date\)\s*(\d{1,2})\s+tháng\s*\(month\)\s*(\d{1,2})\s+năm\s*\(year\)\s*(\d{4})', text)
            invoice_date = f"{int(date_match.group(1)):02d}-{int(date_match.group(2)):02d}-{date_match.group(3)}" if date_match else ""

            buyer_match = re.search(r'Họ tên người mua hàng\s*\(Buyer\):\s+(.*?)\s+Tên đơn vị', text, re.DOTALL)
            buyer_name = buyer_match.group(1).strip().replace("\n", " ") if buyer_match else ""

            vat_match = re.search(r'Thuế suất GTGT.*?:\s+(\d+)%.*?Tiền thuế GTGT.*?:\s+([\d.]+)', text)
            vat_percent = vat_match.group(1) + "%" if vat_match else ""
            vat_amount = vat_match.group(2).replace(".", "") if vat_match else ""

            total_match = re.search(r'Cộng tiền hàng.*?:\s+([\d.]+)', text)
            total_before_tax = total_match.group(1).replace(".", "") if total_match else ""

            total_with_vat = re.search(r'Tổng cộng tiền thanh toán.*?:\s+([\d.]+)', text)
            total_after_tax = total_with_vat.group(1).replace(".", "") if total_with_vat else ""

            # ==== Products ====
            product_rows = []
            for t in tables:
                for row in t:
                    if (
                        row and len(row) >= 7 and
                        isinstance(row[1], str) and
                        re.match(r"\d{13}", row[1])
                    ):
                        product_rows.append(row)

            if not product_rows:
                logging.warning(f"⚠️ No products found in {filename}")
                continue

            for i, row in enumerate(product_rows):
                product_name = row[2].replace("\n", " ").strip()

                if SKIP_FREE_GIFTS and "hàng tặng không bán" in product_name.lower():
                    logging.info(f"⏭️ Skipping gift row: {product_name}")
                    continue

                try:
                    product_code = row[1]
                    unit = row[3]
                    quantity = float(str(row[4]).replace(",", "."))
                    unit_price = float(str(row[5]).replace(",", "").replace(".", ""))
                    total = float(str(row[6]).replace(",", "").replace(".", ""))
                except Exception as e:
                    logging.error(f"❌ Error parsing product row {i}: {row} → {e}")
                    continue

                all_rows.append({
                    "Ngày đặt hàng": ORDER_DATE_DEFAULT if i == 0 else "",
                    "Ngày hóa đơn": invoice_date if i == 0 else "",
                    "Số hóa đơn": invoice_number if i == 0 else "",
                    "Tên khách hàng": buyer_name if i == 0 else "",
                    "Tên đơn vị": "-" if i == 0 else "",
                    "Địa chỉ": "-" if i == 0 else "",
                    "Mã sản phẩm": product_code,
                    "Tên sản phẩm": product_name,
                    "Đơn vị tính": unit,
                    "Số lượng": quantity,
                    "Đơn giá bán (vnđ)": unit_price,
                    "Thành tiền trước thuế (vnđ)": total_before_tax if i == 0 else "",
                    "VAT %": vat_percent if i == 0 else "",
                    "Tiền thuế (vnđ)": vat_amount if i == 0 else "",
                    "Thành tiền sau thuế (vnđ)": total_after_tax if i == 0 else "",
                    "Ghi chú": "",  # Will fill with formula later
                    "Đơn giá bán (HD)": unit_price
                })

        except Exception as e:
            logging.error(f"💥 Failed to process {filename}: {e}")

# ========================
#  Write to Excel
# ========================
df = pd.DataFrame(all_rows)
df.to_excel(OUTPUT_FILE, index=False)

# ========================
#  Add Formulas + Merge Rows
# ========================
wb = load_workbook(OUTPUT_FILE)
ws = wb.active

# Fill Ghi chú formulas
col_q = [cell.value for cell in ws[1]].index("Đơn giá bán (HD)") + 1
col_k = [cell.value for cell in ws[1]].index("Đơn giá bán (vnđ)") + 1
ghi_chu_col = [cell.value for cell in ws[1]].index("Ghi chú") + 1

for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=ghi_chu_col).value = f"={get_column_letter(col_q)}{r}={get_column_letter(col_k)}{r}"

# Merge columns
merge_columns = [
    "Ngày đặt hàng",
    "Ngày hóa đơn",
    "Số hóa đơn",
    "Tên khách hàng",
    "Tên đơn vị",
    "Địa chỉ",
    "Thành tiền trước thuế (vnđ)",
    "VAT %",
    "Tiền thuế (vnđ)",
    "Thành tiền sau thuế (vnđ)"
]
header = [cell.value for cell in ws[1]]
merge_col_indexes = {col: header.index(col) + 1 for col in merge_columns}

row = 2
while row <= ws.max_row:
    start_row = row
    invoice_val = ws.cell(row=row, column=merge_col_indexes["Số hóa đơn"]).value
    end_row = start_row
    while (
        end_row + 1 <= ws.max_row and
        ws.cell(row=end_row + 1, column=merge_col_indexes["Số hóa đơn"]).value in [None, "", invoice_val]
    ):
        end_row += 1

    if end_row > start_row:
        for col_idx in merge_col_indexes.values():
            ws.merge_cells(start_row=start_row, start_column=col_idx, end_row=end_row, end_column=col_idx)

    row = end_row + 1

wb.save(OUTPUT_FILE)
logging.info(f"✅ All done! Saved to {OUTPUT_FILE}")
